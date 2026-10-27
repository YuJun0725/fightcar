#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""整理光电、距离、灰度实验记录；原表只读，结果另存。"""

import csv
from copy import copy, deepcopy
import hashlib
import importlib.util
import json
import re
from pathlib import Path
from statistics import mean, stdev

import numpy as np
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.ticker import MultipleLocator
from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, Reference, ScatterChart, Series
from openpyxl.comments import Comment
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.formula.tokenizer import Tokenizer

HERE = Path(__file__).resolve().parent
INPUT = HERE.parent
FIGURES = HERE / "figures"
SOURCE = INPUT / "光电、灰度.xlsx"
OFFSET = 3.5
BLUE, TEAL, ORANGE, PURPLE = "#2563eb", "#0f9188", "#d97706", "#8b5cf6"
NAVY = "17324D"
FIGURE_INFO = []


def init_sheet(book, name, headers, subtitle):
    sheet = book.create_sheet(name)
    sheet.cell(1, 1, name).font = Font(name="Arial", size=18, bold=True, color=NAVY)
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(5, len(headers)))
    sheet.cell(2, 1, subtitle)
    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=max(5, len(headers)))
    sheet.row_dimensions[1].height = 32
    sheet.row_dimensions[2].height = 38
    sheet.cell(2, 1).alignment = Alignment(wrap_text=True, vertical="center")
    for col, title in enumerate(headers, 1):
        cell = sheet.cell(4, col, title)
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.font = Font(name="Arial", bold=True, color="FFFFFF", size=10)
        cell.alignment = Alignment(wrap_text=True, vertical="center")
    sheet.row_dimensions[4].height = 32
    sheet.freeze_panes = "B5"
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    return sheet


def finish_table(sheet, last, widths=None):
    for row in sheet.iter_rows(min_row=5, max_row=last, max_col=sheet.max_column):
        for cell in row:
            cell.font = Font(name="Arial", size=10, color="008000" if cell.data_type == "f" and "原表快照" in cell.value else "111827")
            cell.alignment = Alignment(vertical="center", wrap_text=True)
            if cell.row % 2:
                cell.fill = PatternFill("solid", fgColor="F1F5F9")
            if cell.data_type == "f" or isinstance(cell.value, (float, int)):
                cell.number_format = "0.00"
        sheet.row_dimensions[row[0].row].height = 24
    for col in range(1, sheet.max_column + 1):
        sheet.column_dimensions[get_column_letter(col)].width = 18
    for col, width in (widths or {}).items():
        sheet.column_dimensions[col].width = width
    sheet.auto_filter.ref = f"A4:{get_column_letter(sheet.max_column)}{last}"


def scatter_chart(sheet, x_col, series_cols, first, last, position, title, x_title, y_title):
    chart = ScatterChart()
    chart.title, chart.x_axis.title, chart.y_axis.title = title, x_title, y_title
    chart.style = 13
    chart.width, chart.height = 22, 12
    chart.scatterStyle = "lineMarker"
    for col in series_cols:
        series = Series(Reference(sheet, min_col=col, min_row=first - 1, max_row=last),
                        Reference(sheet, min_col=x_col, min_row=first, max_row=last), title_from_data=True)
        series.marker.symbol = "circle"
        series.marker.size = 4
        series.graphicalProperties.line.width = 18000
        chart.series.append(series)
    sheet.add_chart(chart, position)


def save_figure(fig, filename, title, note):
    fig.suptitle(title, fontsize=15, fontweight="bold", y=.98)
    fig.text(.5, .015, note, ha="center", fontsize=8.5, color="#475569")
    fig.tight_layout(rect=(0, .055, 1, .92))
    fig.savefig(FIGURES / filename, dpi=190, facecolor="white")
    plt.close(fig)
    FIGURE_INFO.append(dict(file="figures/" + filename, title=title, note=note))


def style_axes(axis, xlabel, ylabel):
    axis.set(xlabel=xlabel, ylabel=ylabel)
    axis.grid(True, color="#e2e8f0", linewidth=.7)
    axis.set_axisbelow(True)


def collect_records(source):
    records = []
    def add(kind, experiment, condition, cell, replicate="", direction="", distance="", angle="", gray="", light="未记录"):
        raw = source[cell].value
        optical = kind == "光电"
        records.append(dict(record_id=f"R{len(records) + 1:03d}", sensor_type=kind, experiment=experiment,
                            condition=condition, replicate=replicate, direction=direction,
                            distance_cm=distance, angle_deg=angle, gray_card=gray, illumination=light,
                            raw_value=raw, raw_unit="尺读数_cm" if optical else "ADC",
                            value=raw - OFFSET if optical else raw, unit="探测距离_cm" if optical else "ADC",
                            source_file=SOURCE.name, source_sheet=source.title, source_cell=cell))
    for surface, start in (("粗糙", 3), ("光面", 6)):
        for trial, row in enumerate(range(start, start + 3), 1):
            add("光电", "P_表面", surface, f"B{row}", trial, "去程")
            add("光电", "P_表面", surface, f"C{row}", trial, "回程")
    for row in range(10, 14):
        for trial, col in enumerate("BCD", 1):
            add("光电", "P_角度", f"{source.cell(row, 1).value}°", f"{col}{row}", trial, angle=source.cell(row, 1).value)
    for row, experiment, condition, angle in ((18, "D_静态", "静态基准（颜色/夹角未注明）", ""),
            (22, "D_颜色", "白色", ""), (23, "D_颜色", "黑色", ""),
            (26, "D_角度", "15°", 15), (27, "D_角度", "30°", 30), (28, "D_角度", "45°", 45)):
        for col in range(2, 10):
            add("距离", experiment, condition, f"{get_column_letter(col)}{row}",
                distance=source.cell(17, col).value, angle=angle)
    for row, distance, light in ((35, 1, "未记录"), (33, 3, "未记录"), (34, 5, "未记录"), (32, 15, "未记录"),
                                 (37, 3, "暗"), (38, 3, "亮")):
        for gray, col in enumerate(range(3, 12), 1):
            add("灰度", "G_距离与色卡" if light == "未记录" else "G_光照", light, f"{get_column_letter(col)}{row}",
                distance=distance, gray=gray, light=light)
    assert len(records) == 126
    return records


def plots(source):
    surface = {}
    fig, axis = plt.subplots(figsize=(9.2, 6))
    for index, (label, start) in enumerate((("粗糙", 3), ("光面", 6))):
        first = [source.cell(row, 2).value - OFFSET for row in range(start, start + 3)]
        back = [source.cell(row, 3).value - OFFSET for row in range(start, start + 3)]
        surface[label] = dict(distance_mean=mean(first), distance_sd=stdev(first), return_mean=mean(back),
                              return_sd=stdev(back), mean_absolute_gap=mean(abs(a - b) for a, b in zip(first, back)))
        for shift, values, color, legend in ((-.17, first, BLUE, "去程"), (.17, back, TEAL, "回程")):
            axis.bar(index + shift, mean(values), width=.3, color=color, alpha=.75,
                     yerr=stdev(values), capsize=5, label=legend if index == 0 else None)
            axis.scatter([index + shift - .045, index + shift, index + shift + .045], values, color="#172554", s=20, zorder=3)
    axis.set_xticks([0, 1], ["粗糙", "光面"])
    style_axes(axis, "目标表面状态", "实际探测距离（cm）")
    axis.set_ylim(0, 65);axis.legend()
    save_figure(fig, "01_photo_surface.png", "光电传感器：不同表面的去程与回程", "去程、回程各 n=3；柱为均值，误差棒为样本标准差，散点为逐次测量；距离均已扣除 3.5 cm。")

    angles = [15, 30, 45, 60]
    triplets = [[source.cell(row, col).value - OFFSET for col in (2, 3, 4)] for row in range(10, 14)]
    fig, axis = plt.subplots(figsize=(9.2, 6))
    axis.errorbar(angles, [mean(v) for v in triplets], yerr=[stdev(v) for v in triplets],
                  color=BLUE, marker="o", capsize=5, label="均值 ± 样本标准差")
    for angle, values in zip(angles, triplets):
        axis.scatter([angle - .7, angle, angle + .7], values, color=TEAL, s=22, zorder=3)
    axis.set_xticks(angles);axis.legend()
    style_axes(axis, "目标夹角（°，沿用原记录）", "实际探测距离（cm）")
    save_figure(fig, "02_photo_angle.png", "光电传感器：目标夹角对探测距离的影响", "每个角度重复测量三次；散点横向微错位仅为显示；距离均已扣除 3.5 cm，未改变角度。")

    distance = np.array([source.cell(17, col).value for col in range(2, 10)], dtype=float)
    baseline = np.array([source.cell(18, col).value for col in range(2, 10)], dtype=float)
    fig, axis = plt.subplots(figsize=(9.2, 6))
    axis.plot(distance, baseline, marker="o", color=BLUE)
    style_axes(axis, "间距（cm）", "记录 ADC 值")
    save_figure(fig, "03_distance_static.png", "距离传感器：静态基准 ADC—距离", "原表速度=0，8 个距离点；基准目标的颜色、夹角和每点采样数未记录；无滤波或跨距离平均。")

    white = np.array([source.cell(22, col).value for col in range(2, 10)], dtype=float)
    black = np.array([source.cell(23, col).value for col in range(2, 10)], dtype=float)
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.5))
    for values, label, color, dash in ((white, "白色", BLUE, "-"), (black, "黑色", TEAL, "-"),
                                        (baseline, "静态基准（颜色未注）", "#64748b", "--")):
        axes[0].plot(distance, values, marker="o", color=color, linestyle=dash, label=label)
    axes[0].legend(fontsize=9)
    delta = (black - white) / white * 100
    axes[1].bar(distance, delta, width=5, color=TEAL)
    axes[1].axhline(0, color="#475569", linewidth=.8)
    style_axes(axes[0], "间距（cm）", "记录 ADC 值")
    style_axes(axes[1], "间距（cm）", "（黑色−白色）/ 白色（%）")
    save_figure(fig, "04_distance_color.png", "距离传感器：不同目标颜色", "黑白各为原表记录点，未知是否为多次均值；未注明颜色的基准不擅自归为本色；差值不代表统计显著性。")

    fig, axis = plt.subplots(figsize=(9.2, 6))
    for row, angle, color in ((26, 15, BLUE), (27, 30, TEAL), (28, 45, ORANGE)):
        axis.plot(distance, [source.cell(row, col).value for col in range(2, 10)], marker="o", color=color, label=f"{angle}°")
    axis.plot(distance, baseline, linestyle="--", color="#64748b", label="静态基准（夹角未注）")
    style_axes(axis, "间距（cm）", "记录 ADC 值");axis.legend()
    save_figure(fig, "05_distance_angle.png", "距离传感器：不同目标夹角", "保留 15°、30°、45° 的全部原表数值；基准夹角未填写，未将其直接标成 0°。")

    grays = np.arange(1, 10)
    gray_rows = [(1, 35), (3, 33), (5, 34), (15, 32)]
    gray_matrix = np.array([[source.cell(row, col).value for col in range(3, 12)] for _, row in gray_rows], float)
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.5))
    for (gap, _), values, color in zip(gray_rows, gray_matrix, (BLUE, TEAL, ORANGE, PURPLE)):
        axes[0].plot(grays, values, marker="o", color=color, label=f"{gap} cm")
    axes[0].legend();axes[0].set_xticks(grays)
    style_axes(axes[0], "色卡编号（1 最白 → 9 最黑）", "记录 ADC 值")
    image = axes[1].imshow(gray_matrix, aspect="auto", cmap="Blues", vmin=0, vmax=2500)
    axes[1].set_xticks(np.arange(9), grays);axes[1].set_yticks(np.arange(4), [f"{gap} cm" for gap, _ in gray_rows])
    axes[1].set(xlabel="色卡编号", ylabel="被测物距离")
    for row in range(4):
        for col in range(9):
            value = gray_matrix[row, col]
            axes[1].text(col, row, f"{value:g}", ha="center", va="center", fontsize=8, color="white" if value > 1300 else "#172554")
    fig.colorbar(image, ax=axes[1], fraction=.04, label="ADC")
    save_figure(fig, "06_gray_distance_cards.png", "灰度传感器：距离与色卡编号", "每个组合仅有一个记录值；1 最白、9 最黑是顺序编号，不是真实反射率；此组环境光未注明。")

    dark = np.array([source.cell(37, col).value for col in range(3, 12)], float)
    bright = np.array([source.cell(38, col).value for col in range(3, 12)], float)
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.5))
    axes[0].plot(grays, dark, color=PURPLE, marker="o", label="暗")
    axes[0].plot(grays, bright, color=ORANGE, marker="o", label="亮")
    axes[0].legend()
    axes[1].bar(grays, bright - dark, color=[BLUE if v >= 0 else ORANGE for v in bright - dark])
    axes[1].axhline(0, color="#475569", linewidth=.8)
    for axis in axes:axis.set_xticks(grays)
    style_axes(axes[0], "色卡编号（1 最白 → 9 最黑）", "记录 ADC 值")
    style_axes(axes[1], "色卡编号", "亮−暗（ADC）")
    save_figure(fig, "07_gray_light.png", "灰度传感器：3 cm 时的环境光影响", "暗、亮均为原表记录值；没有照度 lx 或重复测量，不能据此计算显著性；负差值照实保留。")

    k, b = np.polyfit(1 / baseline, distance, 1)
    prediction = k / baseline + b
    residual = prediction - distance
    fig, axes = plt.subplots(1, 2, figsize=(12.5, 5.5))
    axes[0].scatter(baseline, distance, color=BLUE, label="8 个基准点")
    adc_axis = np.linspace(min(baseline), max(baseline), 300)
    axes[0].plot(adc_axis, k / adc_axis + b, color=TEAL, label="d = a / ADC + b")
    axes[0].legend()
    axes[1].bar(distance, residual, width=5, color=TEAL)
    axes[1].axhline(0, color="#475569", linewidth=.8)
    style_axes(axes[0], "ADC", "间距（cm）")
    style_axes(axes[1], "基准间距（cm）", "同点拟合残差（cm）")
    save_figure(fig, "08_distance_calibration.png", "距离传感器：候选标定与同点拟合残差", "仅用 8 个静态基准点拟合；残差来自同一批拟合点，不是独立验证精度；不得外推或替代实测标定验收。")
    return dict(surface=surface, optical_angle=[dict(angle=a, mean=mean(v), sd=stdev(v)) for a, v in zip(angles, triplets)],
                calibration=dict(method="distance_cm = a / ADC + b", a=float(k), b=float(b),
                    adc_range=[float(min(baseline)), float(max(baseline))], distance_range_cm=[10, 80],
                    fit_points=8, in_sample_rmse_cm=float(np.sqrt(np.mean(residual ** 2))),
                    in_sample_max_abs_cm=float(max(abs(residual))), independent_validation=False),
                black_minus_white_pct=delta.tolist(), bright_minus_dark=(bright - dark).tolist())


def group_sensor_sheets(book):
    layout = [
        ("光电表面", "光电传感器", 4, "I8"),
        ("光电角度", "光电传感器", 27, "I31"),
        ("距离静态", "距离传感器", 4, "G8"),
        ("距离颜色", "距离传感器", 25, "H29"),
        ("距离角度", "距离传感器", 46, "G50"),
        ("距离候选标定", "距离传感器", 67, "N71"),
        ("灰度距离", "灰度传感器", 4, "G8"),
        ("灰度光照", "灰度传感器", 26, "G30"),
    ]
    mapping = {source: (target, shift) for source, target, shift, _ in layout}

    def remap_range(value, local):
        source, address = value.rsplit("!", 1) if "!" in value else (local, value)
        source = source.strip("'")
        if source not in mapping or not re.fullmatch(r"\$?[A-Z]+\$?\d+(?::\$?[A-Z]+\$?\d+)?", address):
            return value
        target, shift = mapping[source]
        address = re.sub(r"(\$?[A-Z]+\$?)(\d+)", lambda m: m[1] + str(int(m[2]) + shift), address)
        return f"'{target}'!{address}"

    def remap_formula(value, local):
        return "=" + "".join(remap_range(token.value, local) if token.subtype == "RANGE" else token.value
                              for token in Tokenizer(value).items)

    targets = {}
    for name in ("光电传感器", "距离传感器", "灰度传感器"):
        sheet = book.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells("A1:G1");sheet["A1"] = name + "实验记录"
        sheet["A1"].font = Font(name="Arial", size=18, bold=True, color=NAVY)
        sheet.merge_cells("A2:G2")
        sheet["A2"] = "按实验内容分区；光电距离均已扣除3.5cm，灰度1最白、9最黑。"
        sheet["A3"] = "测量记录数"
        sheet.merge_cells("C3:G3");sheet["C3"] = "来源：光电、灰度.xlsx / Sheet1"
        sheet.freeze_panes = "B9"
        for col in range(1,27):sheet.column_dimensions[get_column_letter(col)].width=18
        for col,width in {"A":21,"F":22,"G":28,"K":29,"L":23}.items():sheet.column_dimensions[col].width=width
        targets[name]=sheet
    for source_name,target_name,shift,anchor in layout:
        source,target=book[source_name],targets[target_name]
        for row in source:
            for cell in row:
                dest=target.cell(cell.row+shift,cell.column)
                dest.value=remap_formula(cell.value,source_name) if cell.data_type=="f" else cell.value
                if cell.has_style:dest._style=copy(cell._style)
        for merged in source.merged_cells.ranges:
            target.merge_cells(start_row=merged.min_row+shift,start_column=merged.min_col,
                               end_row=merged.max_row+shift,end_column=merged.max_col)
        for index,dim in source.row_dimensions.items():
            if dim.height:target.row_dimensions[index+shift].height=dim.height
        for original in source._charts:
            chart=deepcopy(original)
            for series in chart.series:
                for field in ("val","cat","xVal","yVal","zVal","tx"):
                    obj=getattr(series,field,None)
                    if obj is None:continue
                    for field_ref in ("numRef","strRef"):
                        reference=getattr(obj,field_ref,None)
                        if reference is not None:reference.f=remap_range(reference.f,source_name)
            target.add_chart(chart,anchor)
    counts={"光电传感器":"=COUNT(C9:D14,B32:D35)","距离传感器":"=COUNT(B9:B16,C30:D37,C51:E58)","灰度传感器":"=COUNT(B9:E17,B31:C39)"}
    notes={"光电传感器":(49,"去程/回程各三次，角度各三次；仅记录已扣除3.5cm的实际距离。"),
           "距离传感器":(90,"静态基准颜色/夹角未注。候选标定非独立验证；四速度动态原图和滤波图另存于相邻目录。"),
           "灰度传感器":(49,"色卡编号不是反射率。暗亮在3cm测量；照度、重复采样及固定条件噪声记录仍待补充。")}
    for name,sheet in targets.items():
        sheet["B3"]=counts[name];sheet["B3"].number_format="0"
        row,note=notes[name];sheet.merge_cells(start_row=row,start_column=1,end_row=row+2,end_column=7)
        sheet.cell(row,1,note).alignment=Alignment(wrap_text=True,vertical="center")
    for sheet in list(book):
        if sheet.title not in targets:book.remove(sheet)
    book.active=0


def make_workbook(source_book, records, result):
    def ref(cell):
        return source_book.active[cell].value
    book = Workbook();book.remove(book.active)
    overview = init_sheet(book, "实验总览", ["项目", "当前记录", "状态", "需补充内容"], "原表只读；光电统一减去 3.5 cm；灰度 1 最白、9 最黑；空缺条件明确保留。")
    overview["A5"],overview["B5"] = "光电已扣除零点（cm）", OFFSET
    overview["A6"],overview["B6"] = "色卡顺序", "1 最白，9 最黑（非反射率）"
    overview["A7"],overview["B7"] = "标量观测记录数", "=COUNT('光电表面'!C5:D10,'光电角度'!B5:D8,'距离静态'!B5:B12,'距离颜色'!C5:D12,'距离角度'!C5:E12,'灰度距离'!B5:E13,'灰度光照'!B5:C13)"
    overview["A8"],overview["B8"] = "原表来源", SOURCE.name + " / Sheet1"
    requirements = [
        ("光电：反光条件", "粗糙/光面，各三次去程与回程", "已记录", "目标材质/颜色、光照与旋钮设置"),
        ("光电：角度", "15/30/45/60°，各三次", "已记录", "实际角度定义、目标材质和基准角度"),
        ("距离：静态", "10～80 cm，8 点，速度=0", "已记录", "基准目标颜色/夹角与每点采样量"),
        ("距离：动态", "25/50/100/200 圈/分四份 CSV", "已有原图和滤波图", "时钟/位置对应、重复试验、目标条件"),
        ("距离：白/黑/本色", "白色、黑色各8点；另有未注明颜色的基准", "部分记录", "明确本色是否就是基准；补充材质及重复"),
        ("距离：角度", "15/30/45°，各8点", "已记录", "原表没有明确0°基准、颜色/照明条件"),
        ("传感器标定", "距离8点候选标定；光电零点扣除", "初步整理", "各实物编号与独立验证点；灰度反射率基准"),
        ("噪声处理", "动态 CSV 已做离线5点中值", "部分完成", "固定距离连续原始数据；光电/灰度时序"),
        ("灰度：距离/色卡", "1/3/5/15 cm × 9色卡，共36值", "已记录", "环境光、色卡反射率与每条件重复"),
        ("灰度：环境光", "3 cm × 暗/亮 × 9色卡，共18值", "已记录", "照度lx、重复采样、光源位置"),
    ]
    for row, values in enumerate(requirements, 11):
        for col, value in enumerate(values, 1):overview.cell(row,col,value)
    overview["A23"] = "静态表没有时间序列，不能对相邻距离点直接做时间滤波，也不能由单个记录值计算噪声标准差。"
    overview.merge_cells("A23:D24");overview["A23"].alignment=Alignment(wrap_text=True,vertical="center")
    finish_table(overview,24,{"A":25,"B":58,"C":22,"D":62})
    for row in range(11,21):overview.row_dimensions[row].height=39

    surface = init_sheet(book,"光电表面",["表面","重复","去程距离(cm)","回程距离(cm)","回程−去程(cm)","绝对差(cm)","原表位置"],"原表“距离”列为去程；本表仅列已扣除3.5 cm后的实际距离。")
    for out,(label,start) in enumerate((("粗糙",3),("光面",6))):
        for trial,src in enumerate(range(start,start+3),1):
            row=5+out*3+trial-1
            values=[label,trial,ref(f"B{src}")-OFFSET,ref(f"C{src}")-OFFSET,f"=D{row}-C{row}",f"=ABS(E{row})",f"Sheet1!B{src}:C{src}"]
            for col,value in enumerate(values,1):surface.cell(row,col,value)
    for col,title in enumerate(["表面","去程均值cm","去程样本SD","回程均值cm","回程样本SD","平均绝对差cm","重复数"],1):surface.cell(14,col,title)
    for row,label,first,last in ((15,"粗糙",5,7),(16,"光面",8,10)):
        for col,value in enumerate([label,f"=AVERAGE(C{first}:C{last})",f"=STDEV(C{first}:C{last})",f"=AVERAGE(D{first}:D{last})",f"=STDEV(D{first}:D{last})",f"=AVERAGE(F{first}:F{last})",f"=COUNT(C{first}:C{last})"],1):surface.cell(row,col,value)
    finish_table(surface,16,{"G":28})
    surface.auto_filter.ref="A4:G10"
    for row in range(5,11):surface.cell(row,2).number_format="0"
    for row in (15,16):surface.cell(row,7).number_format="0"
    chart=BarChart();chart.title="表面条件与平均探测距离";chart.y_axis.title="cm";chart.style=10;chart.width=22;chart.height=12
    for col in (2,4):chart.add_data(Reference(surface,min_col=col,min_row=14,max_row=16),titles_from_data=True)
    chart.set_categories(Reference(surface,min_col=1,min_row=15,max_row=16));surface.add_chart(chart,"K4")

    angle=init_sheet(book,"光电角度",["角度(°)","第1次测量(cm)","第2次测量(cm)","第3次测量(cm)","均值(cm)","样本标准差(cm)","原表位置"],"每角度重复测量三次；以下均为已扣除3.5 cm后的实际探测距离。")
    for row,src in enumerate(range(10,14),5):
        values=[ref(f"A{src}"),ref(f"B{src}")-OFFSET,ref(f"C{src}")-OFFSET,ref(f"D{src}")-OFFSET,f"=AVERAGE(B{row}:D{row})",f"=STDEV(B{row}:D{row})",f"Sheet1!A{src}:D{src}"]
        for col,value in enumerate(values,1):angle.cell(row,col,value)
    finish_table(angle,8,{"G":28});scatter_chart(angle,1,[5],5,8,"I4","夹角与平均探测距离","°","cm")

    baseline=init_sheet(book,"距离静态",["间距cm","静态基准ADC","原表位置"],"速度=0；颜色、夹角、照明及每点采样量未记录。小数照原样保存，未假定它是均值。")
    for row,col in enumerate(range(2,10),5):
        letter=get_column_letter(col)
        for index,value in enumerate([ref(f"{letter}17"),ref(f"{letter}18"),f"Sheet1!{letter}17,{letter}18"],1):baseline.cell(row,index,value)
    finish_table(baseline,12,{"C":30});scatter_chart(baseline,1,[2],5,12,"E4","静态基准曲线","cm","ADC")

    color=init_sheet(book,"距离颜色",["间距cm","基准_颜色未注","白色ADC","黑色ADC","黑−白ADC","相对白色差异"],"未知颜色的基准不擅自称为本色；黑白差异不是测量误差或统计显著性。")
    ang=init_sheet(book,"距离角度",["间距cm","基准_夹角未注","15° ADC","30° ADC","45° ADC"],"基准行未注明角度，不能直接等同0°；各距离逐点记录，不平滑。")
    for row,col in enumerate(range(2,10),5):
        letter=get_column_letter(col)
        for index,value in enumerate([f"='距离静态'!A{row}",f"='距离静态'!B{row}",ref(f"{letter}22"),ref(f"{letter}23"),f"=D{row}-C{row}",f'=IFERROR(E{row}/C{row},"")'],1):color.cell(row,index,value)
        for index,value in enumerate([f"='距离静态'!A{row}",f"='距离静态'!B{row}",ref(f"{letter}26"),ref(f"{letter}27"),ref(f"{letter}28")],1):ang.cell(row,index,value)
    finish_table(color,12);finish_table(ang,12)
    for row in range(5,13):color.cell(row,6).number_format="0.00%"
    scatter_chart(color,1,[2,3,4],5,12,"H4","颜色对ADC的影响","cm","ADC")
    scatter_chart(ang,1,[2,3,4,5],5,12,"G4","目标夹角对ADC的影响","cm","ADC")

    gray=init_sheet(book,"灰度距离",["色卡编号","1cm ADC","3cm ADC","5cm ADC","15cm ADC"],"1最白、9最黑；编号不是反射率。四个距离组的原始环境光未填写。")
    light=init_sheet(book,"灰度光照",["色卡编号","暗_3cm ADC","亮_3cm ADC","亮−暗ADC","相对暗差异"],"暗亮为原始条件名，未记录照度lx；负差值照实保留。每组合单个记录值，不计算噪声或显著性。")
    for row,col in enumerate(range(3,12),5):
        letter=get_column_letter(col)
        for index,value in enumerate([ref(f"{letter}31")]+[ref(f"{letter}{r}") for r in (35,33,34,32)],1):gray.cell(row,index,value)
        for index,value in enumerate([ref(f"{letter}31"),ref(f"{letter}37"),ref(f"{letter}38"),f"=C{row}-B{row}",f'=IFERROR(D{row}/B{row},"")'],1):light.cell(row,index,value)
    finish_table(gray,13);finish_table(light,13)
    for row in range(5,14):light.cell(row,5).number_format="0.00%"
    gray.conditional_formatting.add("B5:E13",ColorScaleRule(start_type="min",start_color="EFF6FF",end_type="max",end_color="2563EB"))
    scatter_chart(gray,1,[2,3,4,5],5,13,"G4","距离与灰度色卡","色卡编号","ADC")
    scatter_chart(light,1,[2,3],5,13,"G4","环境光对灰度测量的影响","色卡编号","ADC")

    calibration=init_sheet(book,"距离候选标定",["基准距离cm","基准ADC","1/ADC","模型距离cm","残差cm","绝对残差cm","残差平方"],"候选模型 d=a/ADC+b，仅用于当前8点记录范围。下列是同点拟合残差，不是独立验证精度。")
    for row in range(5,13):
        for col,value in enumerate([f"='距离静态'!A{row}",f"='距离静态'!B{row}",f"=1/B{row}",f"=$L$5*C{row}+$L$6",f"=D{row}-A{row}",f"=ABS(E{row})",f"=E{row}^2"],1):calibration.cell(row,col,value)
    for row,label,value in [(4,"参数","值"),(5,"a", "=SLOPE(A5:A12,C5:C12)"),(6,"b（cm）","=INTERCEPT(A5:A12,C5:C12)"),
                            (7,"同点拟合RMSE cm","=SQRT(AVERAGE(G5:G12))"),(8,"同点最大绝对残差cm","=MAX(F5:F12)"),(9,"拟合点数","=COUNT(A5:A12)"),
                            (10,"最小有效ADC","=MIN(B5:B12)"),(11,"最大有效ADC","=MAX(B5:B12)"),(12,"独立验证","未提供")]:
        calibration.cell(row,11,label);calibration.cell(row,12,value)
    finish_table(calibration,12,{"K":29,"L":23})
    for row in range(5,13):calibration.cell(row,3).number_format="0.000000"
    scatter_chart(calibration,1,[4],5,12,"N4","候选标定的同点输出","基准距离 cm","模型距离 cm")

    for sheet in book:
        if sheet.title.startswith("光电"):sheet.sheet_properties.tabColor="2563EB"
        elif sheet.title.startswith("距离"):sheet.sheet_properties.tabColor="0F9188"
        elif sheet.title.startswith("灰度"):sheet.sheet_properties.tabColor="D97706"
    book.calculation.fullCalcOnLoad=True
    group_sensor_sheets(book)
    path=HERE/"传感器实验整理记录_按传感器合并.xlsx"
    book.save(path)
    return path


def write_report(result, records):
    lines=["# 传感器实验数据记录与绘图","", "本记录依据老师提出的反光条件、目标夹角、动态测距、颜色、标定/噪声、灰度与光照要求整理。原始工作簿只读保留，新增整理文件与图在本目录。", "", "## 已确认的数据解释", "", "- 光电读数为尺上位置，统一减去传感器末端 3.5 cm；整理表仅显示扣除后的距离，尺上读数保留在源文件中。", "- 光电角度的三列是三次重复测量，计算均值与样本标准差。", "- 灰度色卡编号 1 最白、9 最黑；编号不是实测反射率。", "- 距离与灰度表中的ADC小数/整数照原表保存；未注明每点是单次值还是若干次均值，不擅自假定。", "", "## 记录清单", "", "共整理 126 个有来源单元格的标量记录：光电 24、距离 48、灰度 54。另有四速度动态 CSV，与本表静态记录分开保留。", "", "| 项目 | 条件及数量 | 当前状态 |", "|---|---|---|", "| 光电表面 | 粗糙/光面 × 去程/回程 × 3次 | 有记录及重复性统计 |", "| 光电角度 | 15/30/45/60° × 3次 | 有记录及重复性统计 |", "| 距离静态 | 10～80cm，每10cm一点，速度0 | 有8点，颜色/夹角未注明 |", "| 距离颜色 | 白色/黑色 × 8距离 | 有16值，本色未明确 |", "| 距离角度 | 15/30/45° × 8距离 | 有24值，基准角度未明确 |", "| 灰度距离/色卡 | 1/3/5/15cm × 9色卡 | 有36值，环境光未注明 |", "| 灰度光照 | 暗/亮 × 3cm × 9色卡 | 有18值，照度和重复未记录 |", "", "## 光电记录摘要", "", "所有距离单位cm；±为样本标准差，n=3。", "", "| 表面 | 去程均值±SD | 回程均值±SD | 平均绝对回程差 |", "|---|---:|---:|---:|"]
    for label,stats in result["surface"].items():
        lines.append(f'| {label} | {stats["distance_mean"]:.3f} ± {stats["distance_sd"]:.3f} | {stats["return_mean"]:.3f} ± {stats["return_sd"]:.3f} | {stats["mean_absolute_gap"]:.3f} |')
    lines += ["", "| 角度 | 平均探测距离cm | 样本标准差cm |", "|---|---:|---:|"]
    for row in result["optical_angle"]:lines.append(f'| {row["angle"]}° | {row["mean"]:.3f} | {row["sd"]:.3f} |')
    lines += ["", "## 本批数据的主要观察", "", "- 光面目标的平均探测距离大于粗糙目标；光面回程记录的离散程度也较大。这是当前三次记录的现象，目标材质/颜色未填写，不能单独归因为某一个表面光学参数。", "- 光电夹角从15°增至60°时，平均探测距离由约33.67cm降至24.23cm。", "- 距离传感器的黑白目标差异随距离变化：80cm时黑色247.5、白色275 ADC，相对白色约低10%；10cm时黑色略高于白色，不作统一方向的结论。", "- 灰度记录在更近距离具有更大的色卡数值跨度：1cm时1号与9号差2265 ADC，15cm时差60 ADC；尚不能据此给出最优工作距离或噪声意义上的分辨率。", "- 3cm灰度光照试验中，亮−暗在1号色卡为+504 ADC，在8、9号分别为−14、−27 ADC。负差保留，不能概括为增亮总是使ADC上升。"]
    model=result["calibration"]
    lines += ["", "## 距离候选标定", "", f'使用原表8个静态基准点，拟合 `d(cm) = {model["a"]:.6f} / ADC + {model["b"]:.6f}`。仅在已记录的10～80cm、ADC 276.8～2593.3内讨论，不外推。', "", f'同点拟合RMSE约 {model["in_sample_rmse_cm"]:.3f} cm，最大绝对残差约 {model["in_sample_max_abs_cm"]:.3f} cm。**这不是独立验证精度，也不是传感器的噪声标准差。** 原表基准目标颜色/角度未知，当前仅作候选模型，不能当各传感器均已完成标定。', "", "色卡只有编号，没有真实反射率，因此没有编造灰度反射率标定。光电开关记录的是触发位置及回程，不与连续测距ADC混为同一种标定。", "", "## 噪声处理边界", "", "- 距离动态的四份CSV已另存原图和离线5点中值结果，见 ../raw_adc_distance/ 和 ../filtered_adc_distance/。", "- 本表静态/颜色/角度/灰度多为每条件单个汇总值，没有对应的固定条件连续原始采样，不能计算时序噪声σ、信噪比或滤波改善率。", "- 不把相邻距离或相邻色卡当作同一时间信号来做中值/均值滤波；这样会把真实条件变化混入噪声。", "- 光电三次探测位置的标准差只表示边界重复性，不代表光电IO时间噪声。", "", "## 还需补充的实验记录", "", "1. 各传感器的型号/编号/通道、实验时间；目标材质、尺寸、颜色、粗糙度，夹角定义与照度。", "2. 明确距离静态基准是否为本色、是否为正对；没有明确记录前不作该归类。", "3. 固定距离下连续采集原始ADC/IO，保留时间戳；按相同条件重复，再评价噪声和实时滤波。", "4. 距离标定后另选未参与拟合的距离点做独立验证；灰度需要有反射率依据的色卡才能报告物理标定。", "5. 动态滑台同步位置/时间并重复各速度，保留停止后ADC变化现象；滤波不等于已补偿时间差。", "", "## 图表", ""]
    for item in FIGURE_INFO:
        lines += [f'### {item["title"]}', "", f'![{item["title"]}]({item["file"]})', "", item["note"], ""]
    (HERE/"实验记录与缺项.md").write_text("\n".join(lines),encoding="utf-8")
    cards="\n".join(f'<section><h2>{item["title"]}</h2><img src="{item["file"]}" alt="{item["title"]}"><p>{item["note"]}</p></section>' for item in FIGURE_INFO)
    (HERE/"查看全部图表.html").write_text('<!doctype html><html lang="zh-CN"><meta charset="utf-8"><title>传感器实验图表</title><style>body{font-family:Arial,"Microsoft YaHei",sans-serif;max-width:1200px;margin:32px auto;padding:0 18px;color:#17324d;background:#f5f7fa}section{background:white;margin:24px 0;padding:22px;border-radius:12px}img{width:100%;height:auto}p{line-height:1.7;color:#475569}h1{font-size:30px}h2{font-size:21px}</style><h1>传感器实验记录与图表</h1><p>依据原表整理。光电尺读数减去3.5cm；灰度编号1最白、9最黑。原表保持不变；缺失条件未补造。</p>'+cards+'</html>',encoding="utf-8")


def main():
    digest=hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    source_book=load_workbook(SOURCE,data_only=False)
    source=source_book.active
    spec=importlib.util.spec_from_file_location("plot_style",INPUT/"raw_adc_distance"/"plot_raw.py")
    style=importlib.util.module_from_spec(spec);spec.loader.exec_module(style);style.setup_style()
    FIGURES.mkdir(exist_ok=True)
    records=collect_records(source)
    with (HERE/"全部测量记录.csv").open("w",encoding="utf-8-sig",newline="") as output:
        writer=csv.DictWriter(output,fieldnames=list(records[0]));writer.writeheader();writer.writerows(records)
    result=plots(source)
    workbook=make_workbook(source_book,records,result)
    write_report(result,records)
    manifest=dict(source_file=SOURCE.name,source_sha256=digest,source_sheet=source.title,scalar_records=126,
                  confirmed=dict(optical_offset_cm=OFFSET,optical_angle_repeats=3,gray_1="最白",gray_9="最黑"),
                  summaries=result,figures=FIGURE_INFO)
    (HERE/"整理依据.json").write_text(json.dumps(manifest,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest()==digest,"原工作簿在整理过程中变化，请重新读取"
    print("WORKBOOK",workbook)
    print("RECORDS",len(records),"FIGURES",len(FIGURE_INFO))
    print(json.dumps(result,ensure_ascii=True))


if __name__=="__main__":main()
