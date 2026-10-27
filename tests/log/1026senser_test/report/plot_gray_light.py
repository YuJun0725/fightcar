"""绘制 3 cm 距离下暗、亮两种条件的灰度原始数据。"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from openpyxl import load_workbook

root = Path(__file__).resolve().parent
workbook = load_workbook(root.parent / "光电、灰度.xlsx", data_only=True, read_only=True)
sheet = workbook["Sheet1"]
cards = [sheet.cell(31, col).value for col in range(3, 12)]
dark = [sheet.cell(37, col).value for col in range(3, 12)]
bright = [sheet.cell(38, col).value for col in range(3, 12)]
workbook.close()

plt.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans SC"],
    "axes.unicode_minus": False,
    "pdf.fonttype": 42,
    "font.size": 10,
})
fig, ax = plt.subplots(figsize=(6.4, 3.9))
ax.plot(cards, dark, "o-", color="#2563eb", label="暗")
ax.plot(cards, bright, "s-", color="#d97706", label="亮")
ax.set_xticks(cards)
ax.set_xlabel("色卡编号（1 最白，9 最黑）")
ax.set_ylabel("ADC 值")
ax.set_ylim(bottom=0)
ax.grid(True, color="#dcdcdc", linewidth=0.6)
ax.set_axisbelow(True)
ax.legend(frameon=False)
fig.tight_layout()
fig.savefig(root / "figures" / "fig_gray_light_raw.pdf", bbox_inches="tight")
plt.close(fig)
