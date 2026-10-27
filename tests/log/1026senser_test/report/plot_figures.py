# -*- coding: utf-8 -*-
"""生成实验报告用图（独立于记录目录中的原始图）。

数据来源：
  上级目录 光电、灰度.xlsx          —— 光电、距离静态/颜色/角度、灰度
  ../raw_adc_distance/aligned_raw_samples.csv
  ../filtered_adc_distance/filtered_motion_samples.csv
输出：figures/*.pdf
"""
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
DATA = HERE.parent
FIG = HERE / "figures"
FIG.mkdir(exist_ok=True)

plt.rcParams.update({
    "font.family": "sans-serif",
    "font.sans-serif": ["Microsoft YaHei", "SimHei", "Noto Sans SC"],
    "axes.unicode_minus": False,
    "pdf.fonttype": 42,
    "ps.fonttype": 42,
    "font.size": 9.5,
    "axes.labelsize": 9.5,
    "xtick.labelsize": 8.5,
    "ytick.labelsize": 8.5,
    "legend.fontsize": 8.5,
    "axes.grid": True,
    "grid.color": "#dcdcdc",
    "grid.linewidth": 0.6,
    "axes.edgecolor": "#333333",
    "axes.linewidth": 0.8,
    "lines.linewidth": 1.4,
    "lines.markersize": 4,
    "savefig.bbox": "tight",
    "savefig.pad_inches": 0.03,
})

BLUE, TEAL, ORANGE, PURPLE, GRAY = "#2563eb", "#0f9188", "#d97706", "#8b5cf6", "#6b7280"


def style(ax, xlabel="", ylabel=""):
    ax.set_xlabel(xlabel)
    ax.set_ylabel(ylabel)
    ax.set_axisbelow(True)


def save(fig, name):
    fig.savefig(FIG / name, facecolor="white")
    plt.close(fig)
    print("saved", name)


# ---------------------------------------------------------------- 数据
wb = load_workbook(DATA / "光电、灰度.xlsx", data_only=True)
ws = wb["Sheet1"]
OFFSET = 3.5

surfaces = {
    "粗糙": dict(out=[ws.cell(r, 2).value for r in (3, 4, 5)],
              ret=[ws.cell(r, 3).value for r in (3, 4, 5)]),
    "光面": dict(out=[ws.cell(r, 2).value for r in (6, 7, 8)],
              ret=[ws.cell(r, 3).value for r in (6, 7, 8)]),
}
angles = [ws.cell(r, 1).value for r in (10, 11, 12, 13)]
angle_raw = [[ws.cell(r, c).value for c in (2, 3, 4)] for r in (10, 11, 12, 13)]
dists = [ws.cell(17, c).value for c in range(2, 10)]
d_static = [ws.cell(18, c).value for c in range(2, 10)]
d_white = [ws.cell(22, c).value for c in range(2, 10)]
d_black = [ws.cell(23, c).value for c in range(2, 10)]
d_15 = [ws.cell(26, c).value for c in range(2, 10)]
d_30 = [ws.cell(27, c).value for c in range(2, 10)]
d_45 = [ws.cell(28, c).value for c in range(2, 10)]
cards = [ws.cell(31, c).value for c in range(3, 12)]
gray = {
    "1cm": [ws.cell(35, c).value for c in range(3, 12)],
    "3cm": [ws.cell(33, c).value for c in range(3, 12)],
    "5cm": [ws.cell(34, c).value for c in range(3, 12)],
    "15cm": [ws.cell(32, c).value for c in range(3, 12)],
}
dark = [ws.cell(37, c).value for c in range(3, 12)]
bright = [ws.cell(38, c).value for c in range(3, 12)]

# 拟合 d = a/ADC + b
adc = np.array(d_static, dtype=float)
dist = np.array(dists, dtype=float)
k, b = np.polyfit(1 / adc, dist, 1)
print(f"fit a={k:.6f} b={b:.6f}")
pred = k / adc + b
resid = pred - dist
print("rmse", float(np.sqrt(np.mean(resid ** 2))), "max", float(np.max(np.abs(resid))))

# ---------------------------------------------------------------- 图 1 光电：表面 + 角度
fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.3))
ax = axes[0]
pos = np.arange(2)
w = 0.32
for i, (key, label, color) in enumerate([("out", "去程", BLUE), ("ret", "回程", TEAL)]):
    data = [surfaces[s][key] for s in surfaces]
    means = [float(np.mean(v)) for v in data]
    sds = [float(np.std(v, ddof=1)) for v in data]
    ax.bar(pos + (i - 0.5) * w, means, w, color=color, alpha=0.85, label=label,
           edgecolor="white", linewidth=0.6)
    ax.errorbar(pos + (i - 0.5) * w, means, yerr=sds, fmt="none", ecolor="#111111",
                capsize=3.5, elinewidth=1.1)
    for j, vals in enumerate(data):
        jitter = np.linspace(-0.06, 0.06, len(vals))
        ax.plot(np.full(len(vals), pos[j] + (i - 0.5) * w) + jitter, vals, "o",
                color="#111827", markersize=3.0, zorder=5)
ax.set_xticks(pos, [s + " (n=3)" for s in surfaces])
ax.set_ylim(0, 60)
ax.legend(loc="upper left", frameon=False, ncol=2)
style(ax, "目标表面", "探测距离 / cm")

ax = axes[1]
means = [float(np.mean(v)) for v in angle_raw]
sds = [float(np.std(v, ddof=1)) for v in angle_raw]
ax.plot(angles, means, "-o", color=BLUE, label="均值", zorder=3)
ax.errorbar(angles, means, yerr=sds, fmt="none", ecolor="#111111", capsize=3.5, elinewidth=1.1)
for a, vals in zip(angles, angle_raw):
    jitter = np.linspace(-0.9, 0.9, len(vals))
    ax.plot(np.full(len(vals), a) + jitter, vals, "o", color="#111827", markersize=3.2, zorder=4)
ax.set_xticks(angles)
ax.set_xlim(8, 67)
ax.set_ylim(20, 38)
style(ax, "目标夹角 / (°)", "探测距离 / cm")
fig.tight_layout(w_pad=1.6)
save(fig, "fig_photo.pdf")

# ---------------------------------------------------------------- 图 3 距离静态 + 标定
fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.1))
ax = axes[0]
ax.plot(dists, d_static, "o-", color=BLUE)
ax.set_xlim(0, 90)
ax.set_ylim(0, 2900)
style(ax, "间距 / cm", "ADC 值")

ax = axes[1]
lo, hi = 5, 85
ax.plot([lo, hi], [lo, hi], "--", color=GRAY, linewidth=1.0, label="y = x")
ax.plot(dist, pred, "o", color=ORANGE, markersize=4.5, label="模型距离")
ax.set_xlim(0, 90)
ax.set_ylim(0, 90)
ax.legend(loc="upper left", frameon=False)
style(ax, "实测距离 / cm", "模型距离 / cm")
fig.tight_layout(w_pad=1.6)
save(fig, "fig_distance_static.pdf")

# ---------------------------------------------------------------- 图 4 距离：颜色 / 角度
fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.1))
ax = axes[0]
ax.plot(dists, d_white, "o-", color=BLUE, label="白色")
ax.plot(dists, d_black, "s-", color="#111827", label="黑色")
ax.set_xlim(0, 90)
ax.set_ylim(0, 2900)
ax.legend(loc="upper right", frameon=False)
style(ax, "间距 / cm", "ADC 值")

ax = axes[1]
ax.plot(dists, d_static, "o--", color=GRAY, label="基准（角度未注明）")
ax.plot(dists, d_15, "^-", color=BLUE, label="15°")
ax.plot(dists, d_30, "s-", color=TEAL, label="30°")
ax.plot(dists, d_45, "d-", color=ORANGE, label="45°")
ax.set_xlim(0, 90)
ax.set_ylim(0, 2900)
ax.legend(loc="upper right", frameon=False, ncol=2)
style(ax, "间距 / cm", "ADC 值")
fig.tight_layout(w_pad=1.6)
save(fig, "fig_distance_color_angle.pdf")

# ---------------------------------------------------------------- 图 5 灰度：距离与色卡
fig, ax = plt.subplots(figsize=(6.4, 3.9))
for key, color, marker in [("1cm", BLUE, "o"), ("3cm", TEAL, "s"),
                           ("5cm", ORANGE, "^"), ("15cm", PURPLE, "d")]:
    ax.plot(cards, gray[key], marker=marker, color=color, label=key, markersize=4)
ax.set_xticks(cards)
ax.set_xlim(0.5, 9.5)
ax.set_ylim(0, 2700)
ax.legend(title="测量距离", loc="upper right", frameon=False)
style(ax, "色卡编号（1 最白，9 最黑）", "ADC 值")
save(fig, "fig_gray_cards.pdf")

# ---------------------------------------------------------------- 图 6 灰度：环境光
fig, axes = plt.subplots(1, 2, figsize=(7.2, 3.2))
ax = axes[0]
x = np.arange(len(cards))
w = 0.38
ax.bar(x - w / 2, dark, w, color="#3f4a5a", label="暗")
ax.bar(x + w / 2, bright, w, color=ORANGE, label="亮")
ax.set_xticks(x, cards)
ax.set_ylim(0, 2000)
ax.legend(loc="upper right", frameon=False)
style(ax, "色卡编号（3 cm）", "ADC 值")

ax = axes[1]
delta = [b_ - d_ for b_, d_ in zip(bright, dark)]
colors = [TEAL if v >= 0 else "#c2410c" for v in delta]
ax.bar(x, delta, 0.55, color=colors)
ax.axhline(0, color="#333333", linewidth=0.8)
ax.set_xticks(x, cards)
style(ax, "色卡编号（3 cm）", "亮 − 暗 / ADC")
fig.tight_layout(w_pad=1.8)
save(fig, "fig_gray_light.pdf")

# ---------------------------------------------------------------- 图 7 动态：原始四速度
raw = pd.read_csv(DATA / "raw_adc_distance" / "aligned_raw_samples.csv", encoding="utf-8-sig")
fig, axes = plt.subplots(2, 2, figsize=(7.2, 5.4), sharex=True, sharey=True)
for ax, rpm in zip(axes.flat, [25, 50, 100, 200]):
    part = raw[raw["motor_rpm"] == rpm]
    motion = part[part["phase"] == "motion"].sort_values("estimated_distance_cm")
    still = part[part["phase"] != "motion"]
    ax.plot(still["estimated_distance_cm"], still["adc_value"], "o", color="#b0b7c3",
            markersize=2.4, label="启停前后采样")
    ax.plot(motion["estimated_distance_cm"], motion["adc_value"], "-o", color=BLUE,
            markersize=2.4, linewidth=1.1, label="运动区间")
    ax.set_xlim(0, 90)
    ax.set_ylim(0, 2900)
    ax.set_title(f"{rpm} 圈/分（约 {rpm / 8:g} cm/s，{len(motion)} 点）", fontsize=8.8)
axes[0, 0].legend(loc="upper right", frameon=False)
for ax in axes[1]:
    style(ax, "估算距离 / cm", "原始 ADC 值")
for ax in axes[:, 0]:
    style(ax, "", "原始 ADC 值")
fig.tight_layout(w_pad=1.2, h_pad=1.2)
save(fig, "fig_dynamic_raw.pdf")

# ---------------------------------------------------------------- 图 8 动态：中值滤波
fil = pd.read_csv(DATA / "filtered_adc_distance" / "filtered_motion_samples.csv",
                  encoding="utf-8-sig")
fig, axes = plt.subplots(2, 2, figsize=(7.2, 5.4), sharex=True, sharey=True)
for ax, rpm in zip(axes.flat, [25, 50, 100, 200]):
    part = fil[fil["motor_rpm"] == rpm].sort_values("estimated_distance_cm")
    ax.plot(part["estimated_distance_cm"], part["adc_value"], "-", color="#b0b7c3",
            linewidth=1.0, label="原始值")
    ax.plot(part["estimated_distance_cm"], part["adc_median5"], "-", color=TEAL,
            linewidth=1.4, label="5 点中值")
    ax.set_xlim(0, 90)
    ax.set_ylim(0, 2900)
    ax.set_title(f"{rpm} 圈/分（{len(part)} 点）", fontsize=8.8)
axes[0, 0].legend(loc="upper right", frameon=False)
for ax in axes[1]:
    style(ax, "估算距离 / cm", "ADC 值")
for ax in axes[:, 0]:
    style(ax, "", "ADC 值")
fig.tight_layout(w_pad=1.2, h_pad=1.2)
save(fig, "fig_dynamic_filtered.pdf")

# ---------------------------------------------------------------- 统计输出
print("\n表面统计")
for s in surfaces:
    out = np.array(surfaces[s]["out"], dtype=float) - OFFSET
    ret = np.array(surfaces[s]["ret"], dtype=float) - OFFSET
    print(f"{s}: 去程 {out.mean():.3f}±{out.std(ddof=1):.3f}  回程 {ret.mean():.3f}±{ret.std(ddof=1):.3f}  "
          f"|Δ|均值 {np.mean(np.abs(ret - out)):.3f}")
print("\n角度统计")
for a, vals in zip(angles, angle_raw):
    v = np.array(vals, dtype=float) - OFFSET
    print(f"{a}: {v.mean():.3f}±{v.std(ddof=1):.3f}")
print("\n灰度 1cm/15cm 端点差", gray["1cm"][0] - gray["1cm"][-1], gray["15cm"][0] - gray["15cm"][-1])
print("亮-暗", [round(b_ - d_, 1) for b_, d_ in zip(bright, dark)])
